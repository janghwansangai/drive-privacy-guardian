"""Phase 3 GUI tests (pytest-qt, offscreen). Real AuthManager/AuditRunner over the fakes."""

from __future__ import annotations

import datetime as dt
import os
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QFileDialog, QTableWidget

from dpg.core.auth import AccessLevel, AuthManager
from dpg.core.auth.secret_store import KEY_TOKEN
from dpg.core.drive.client import DriveClient
from dpg.gui import main_window as mw
from dpg.gui.context import AppContext, Prefs
from dpg.gui.main_window import MainWindow, ReviewDialog, SettingsDialog, folder_id_from
from dpg.gui.results_model import COL_NAME, COL_PRIVACY
from dpg.gui.wizard import PAGE_CLIENT, PAGE_GUIDE, PAGE_LOGIN, PAGE_MODE, SetupWizard
from tests.fakes.fake_drive import FakeDrive
from tests.fakes.fake_google_auth import FakeGoogle, MemorySecretStore

pytestmark = pytest.mark.enable_socket  # loopback login in the fake browser

ME = "teacher@school.example"


@dataclass
class Env:
    google: FakeGoogle
    store: MemorySecretStore
    fake: FakeDrive
    opened: list[str] = field(default_factory=list)
    notices: list[tuple[str, str]] = field(default_factory=list)
    ctx: AppContext | None = None


class FakeAutostart:
    """Never touches the real login items of the machine running the tests."""

    def __init__(self) -> None:
        self.on = False
        self.fail: Exception | None = None

    def supported(self) -> bool:
        return True

    def is_enabled(self) -> bool:
        return self.on

    def enable(self) -> None:
        if self.fail is not None:
            raise self.fail
        self.on = True

    def disable(self) -> None:
        self.on = False


@pytest.fixture
def env() -> Env:
    e = Env(FakeGoogle(email=ME), MemorySecretStore(), FakeDrive(me=ME))

    def open_url(url: str) -> bool:
        e.opened.append(url.split("?", 1)[0])
        if url.startswith("https://accounts.google.com/"):
            return e.google.browser(url)
        return True

    e.ctx = AppContext(
        manager_factory=lambda: AuthManager(e.store, e.google.http, e.google.browser),
        service_factory=lambda _m: e.fake.service(),
        open_url=open_url,
        notify=lambda _p, title, text, *a: e.notices.append((title, text)),
        confirm=lambda *a: True,
        prefs=Prefs(),
        autostart=FakeAutostart(),
    )
    return e


def _login(env: Env) -> None:
    assert env.ctx is not None
    env.ctx.manager.import_client(env.google.client_json())
    env.ctx.manager.login(AccessLevel.AUDIT)


def _populate(fake: FakeDrive) -> dict[str, str]:
    ids = {"private": fake.add_file("비공개.txt")}
    ids["link"] = fake.add_file("공지.pdf")
    fake.share(ids["link"], "anyone", "reader")
    ids["ext"] = fake.add_file("외부.xlsx")
    fake.share(ids["ext"], "user", "writer", email="parent@gmail.example")
    ids["evil"] = fake.add_file('=HYPERLINK("http://evil.example","x").xlsx')
    fake.share(ids["evil"], "anyone", "writer")
    other = fake.add_file("남의 파일", owner="other@school.example")
    fake.share(other, "user", "reader", email=ME)
    return ids


def _audit(qtbot: Any, window: MainWindow) -> None:
    qtbot.mouseClick(window.start_btn, mw.Qt.MouseButton.LeftButton)
    qtbot.waitUntil(lambda: window.task is None, timeout=15000)


# --- wizard -----------------------------------------------------------------------------------


def test_wizard_school_flow(qtbot: Any, env: Env, tmp_path: Path) -> None:
    assert env.ctx is not None
    wizard = SetupWizard(env.ctx)
    qtbot.addWidget(wizard)
    wizard.show()
    wizard.mode_page.school.setChecked(True)
    wizard.next()
    assert wizard.currentId() == PAGE_CLIENT
    assert not wizard.client_page.isComplete()

    src = tmp_path / "client_secret_school.json"
    src.write_text(env.google.client_json(), encoding="utf-8")
    wizard.client_page.import_file(src)
    assert wizard.client_page.isComplete()
    assert "키체인" in wizard.client_page.status.text()
    assert not wizard.client_page.trash.isHidden()
    wizard.next()
    assert wizard.currentId() == PAGE_LOGIN

    page = wizard.login_page
    qtbot.mouseClick(page.login_btn, mw.Qt.MouseButton.LeftButton)
    qtbot.waitUntil(lambda: page.done, timeout=15000)
    assert ME in page.status.text()
    assert KEY_TOKEN in env.store.data
    wizard.close()


def test_wizard_rejects_wrong_file(qtbot: Any, env: Env, tmp_path: Path) -> None:
    assert env.ctx is not None
    wizard = SetupWizard(env.ctx)
    qtbot.addWidget(wizard)
    bad = tmp_path / "web.json"
    bad.write_text('{"web": {"client_id": "x"}}', encoding="utf-8")
    wizard.client_page.import_file(bad)
    assert wizard.client_page.status.text().startswith("✗")
    assert "데스크톱 앱" in wizard.client_page.status.text()
    assert not wizard.client_page.isComplete()


def test_wizard_personal_guide(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    wizard = SetupWizard(env.ctx)
    qtbot.addWidget(wizard)
    wizard.show()
    wizard.next()
    assert wizard.currentId() == PAGE_GUIDE
    guide = wizard.guide_page
    assert not guide.isComplete()
    for check in guide.checks:
        check.setChecked(True)
    assert guide.isComplete()
    # the step buttons open console pages in the user's browser
    from PySide6.QtWidgets import QPushButton

    for btn in guide.findChildren(QPushButton):
        if btn.text() == "열기":
            qtbot.mouseClick(btn, mw.Qt.MouseButton.LeftButton)
    assert "https://console.cloud.google.com/apis/library/drive.googleapis.com" in env.opened


def test_wizard_starts_at_login_when_client_exists(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    env.ctx.manager.import_client(env.google.client_json())
    wizard = SetupWizard(env.ctx)
    qtbot.addWidget(wizard)
    assert wizard.startId() == PAGE_LOGIN


def test_wizard_login_page_switches_to_a_new_client(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    env.ctx.manager.import_client(env.google.client_json())
    wizard = SetupWizard(env.ctx)
    qtbot.addWidget(wizard)
    wizard.restart()
    assert "사용 중인 열쇠" in wizard.login_page.key_label.text()
    wizard.login_page.switch_btn.click()
    assert env.ctx.manager.client() is None
    assert wizard.currentId() == PAGE_MODE


def test_switch_client_from_main_window(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert env.ctx is not None
    _login(env)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    assert "열쇠:" in window.account_label.text()
    setups: list[bool] = []
    monkeypatch.setattr(window, "run_setup", lambda: setups.append(True) or False)
    window.switch_btn.click()
    assert setups == [True]
    assert env.ctx.manager.client() is None
    assert not env.ctx.manager.status().logged_in
    assert window.result is None


# --- main window ------------------------------------------------------------------------------


def test_audit_dashboard_filter_export(
    qtbot: Any, env: Env, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert env.ctx is not None
    _login(env)
    ids = _populate(env.fake)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    assert ME in window.account_label.text()

    _audit(qtbot, window)
    assert window.result is not None
    assert "감사 완료" in window.progress_label.text()
    assert window.dash_buttons["all"].text().endswith("4")  # my 4 files (scope: mine)
    assert window.dash_buttons["link"].text().endswith("2")
    assert window.dash_buttons["external"].text().endswith("1")
    assert env.fake.write_calls == []

    window.set_filter("link")
    assert window.proxy.rowCount() == 2
    window.set_filter("all")
    window.search_edit.setText("공지")
    assert window.proxy.rowCount() == 1
    window.search_edit.setText("")
    # "위험도 높은 순" puts the riskiest first (default is folders/files)
    window.sort_combo.setCurrentIndex(1)
    first = window.proxy.index(0, COL_NAME).data()
    assert "HYPERLINK" in first

    # export XLSX (masked) — the formula-looking name must stay a literal string
    out = tmp_path / "report.xlsx"
    monkeypatch.setattr(mw.ExportDialog, "exec", lambda self: 1)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(out), ""))
    window.export_report()
    ws = load_workbook(out).active
    names = [c.value for c in ws["B"][1:]]
    assert all(isinstance(v, str) and not v.startswith("=") for v in names)
    assert not any("parent@gmail" in str(c.value) for row in ws.iter_rows() for c in row)
    if os.name == "posix":
        assert out.stat().st_mode & 0o777 == 0o600
    assert env.notices[-1][0] == "보고서 저장"

    # unmasked CSV
    csv_out = tmp_path / "report.csv"

    def csv_dialog(self: mw.ExportDialog) -> int:
        self.fmt.setCurrentIndex(self.fmt.findData(".csv"))
        self.mask_check.setChecked(False)
        return 1

    monkeypatch.setattr(mw.ExportDialog, "exec", csv_dialog)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(csv_out), ""))
    window.export_report()
    text = csv_out.read_text(encoding="utf-8-sig")
    assert "parent@gmail.example" in text
    assert "'=HYPERLINK" in text
    assert ids["ext"] in text


def test_shared_scope_shows_review_items(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    _login(env)
    _populate(env.fake)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    window.scope_combo.setCurrentIndex(window.scope_combo.findData("shared"))
    _audit(qtbot, window)
    assert window.dash_buttons["review"].text().endswith("1")
    assert "안전하다는 뜻이 아니므로" in window.notice_label.text()


def test_cancel_and_resume(qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch) -> None:
    assert env.ctx is not None
    _login(env)
    for i in range(15):
        folder = env.fake.add_folder(f"반{i}")
        for j in range(100):
            env.fake.add_file(f"f{i}-{j}", parent=folder)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)

    original = DriveClient.list_files_page
    calls = {"n": 0}

    def cancel_after_first_page(self: DriveClient, **kw: Any) -> Any:
        calls["n"] += 1
        if calls["n"] == 2:
            window.cancel_event.set()  # as if the user pressed 중단
        return original(self, **kw)

    monkeypatch.setattr(DriveClient, "list_files_page", cancel_after_first_page)
    _audit(qtbot, window)
    assert window.result is None
    assert "중단" in window.progress_label.text()

    monkeypatch.setattr(DriveClient, "list_files_page", original)
    _audit(qtbot, window)
    assert window.result is not None
    assert len(window.result.items) == 1515
    assert "이어서" in window.progress_label.text()


def test_expired_login_asks_to_log_in_again(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert env.ctx is not None
    _login(env)
    env.google.expire_all_refresh_tokens()
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    setups: list[bool] = []
    monkeypatch.setattr(window, "run_setup", lambda: setups.append(True) or False)
    _audit(qtbot, window)
    assert setups == [True]
    assert any("다시 로그인" in text for _t, text in env.notices)
    assert "로그인되어 있지 않습니다" in window.account_label.text()


def test_folder_scope_requires_input(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    _login(env)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    window.scope_combo.setCurrentIndex(window.scope_combo.findData("folder"))
    window.start_audit()
    assert env.notices[-1][0] == "감사 범위"
    assert window.task is None


def test_folder_id_from_url() -> None:
    assert (
        folder_id_from("https://drive.google.com/drive/folders/1AbC_d-9?usp=sharing") == "1AbC_d-9"
    )
    assert folder_id_from("https://drive.google.com/drive/u/0/folders/1XyZ") == "1XyZ"
    assert folder_id_from("  1RawId  ") == "1RawId"


def test_settings_and_auto_logout(qtbot: Any, env: Env, isolated_app_home: Path) -> None:
    assert env.ctx is not None
    _login(env)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    dialog = SettingsDialog(env.ctx, window)
    dialog.auto_logout.setChecked(True)
    dialog.domains.setText("school.go.kr; Other.Example")
    dialog._save()
    saved = Prefs.load()
    assert saved.auto_logout_on_exit is True
    assert saved.internal_domains == ["school.go.kr", "other.example"]
    if os.name == "posix":
        assert Prefs.path().stat().st_mode & 0o777 == 0o600

    window.close()
    assert KEY_TOKEN not in env.store.data  # logged out on exit
    assert env.google.revoked


def test_no_qt_networking_loaded(qtbot: Any, env: Env) -> None:
    """Principle 1 / D-008: QtNetwork bypasses the Python socket guard, so it must never load."""
    assert env.ctx is not None
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    wizard = SetupWizard(env.ctx)
    qtbot.addWidget(wizard)
    loaded = [m for m in sys.modules if m.startswith(("PySide6.QtNetwork", "PySide6.QtWebEngine"))]
    assert loaded == []


# --- Phase 4: detection in the GUI -----------------------------------------------------------


def test_detection_flow_with_permission_upgrade(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    from dpg.core.auth.scopes import AccessLevel as Level
    from dpg.core.policy import DetectStatus

    assert env.ctx is not None
    _login(env)  # AUDIT only
    fixtures = Path(__file__).resolve().parents[1] / "fixtures" / "synthetic"
    csv_id = env.fake.add_file(
        "6-2_학생_연락처.csv",
        mime_type="text/csv",
        content=(fixtures / "6-2_학생_연락처.csv").read_bytes(),
    )
    env.fake.share(csv_id, "anyone", "reader")
    env.fake.add_file(
        "스캔_가상.pdf",
        mime_type="application/pdf",
        content=(fixtures / "스캔_가상.pdf").read_bytes(),
    )
    dialogs: list[Any] = []
    env.ctx.show_dialog = dialogs.append
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)

    upgrades: list[Level] = []

    def fake_setup(level: Level = Level.AUDIT) -> bool:
        upgrades.append(level)
        env.ctx.manager.login(level)  # type: ignore[union-attr]
        return True

    monkeypatch.setattr(window, "run_setup", fake_setup)
    window.detect_check.setChecked(True)
    _audit(qtbot, window)
    assert upgrades == [Level.DETECT]  # extra scope requested only when first needed
    assert window.detections is not None
    assert window.detections[csv_id].status is DetectStatus.DETECTED
    assert window.dash_buttons["detected"].text().endswith("1")
    assert window.dash_buttons["urgent"].text().endswith("1")
    window.set_filter("urgent")
    assert window.proxy.rowCount() == 1
    assert "주민" in window.proxy.index(0, COL_PRIVACY).data()
    window.set_filter("review")
    assert "검사 불가" in window.proxy.index(0, COL_PRIVACY).data()
    assert env.fake.write_calls == []

    # review: masked, in memory
    window.set_filter("urgent")
    window.table.selectRow(0)
    window.review_selected()
    qtbot.waitUntil(lambda: bool(dialogs), timeout=10000)
    dialog = dialogs[0]
    assert isinstance(dialog, ReviewDialog)
    text = (fixtures / "6-2_학생_연락처.csv").read_text(encoding="utf-8")
    import re

    first_rrn = re.search(r"\d{6}-\d{7}", text).group()  # type: ignore[union-attr]
    from PySide6.QtWidgets import QTableWidget

    table = dialog.findChild(QTableWidget)
    shown = " ".join(table.item(r, 2).text() for r in range(table.rowCount()))
    assert first_rrn not in shown
    assert first_rrn[:8] in shown  # "YYMMDD-G******"

    # exclude
    window.exclude_selected()
    assert window.detections[csv_id].status is DetectStatus.EXCLUDED
    window.detect_check.setChecked(True)
    window.restart_check.setChecked(True)
    _audit(qtbot, window)
    assert window.detections[csv_id].status is DetectStatus.EXCLUDED  # remembered


def test_detection_declined_keeps_audit_only(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    _login(env)
    env.ctx.confirm = lambda *a: False
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    window.detect_check.setChecked(True)
    window.start_audit()
    assert window.task is None  # nothing started without the user's consent


def test_about_shows_hwp_notice(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    window.show_about()
    assert "한컴의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다" in env.notices[-1][1]


# --- Phase 5: permission changes in the GUI --------------------------------------------------


def _writable(env: Env) -> None:
    env.fake.read_only = False


def _select_file(window: MainWindow, name: str) -> None:
    for row in range(window.proxy.rowCount()):
        if window.proxy.index(row, COL_NAME).data() == name:
            window.table.selectRow(row)
            return
    raise AssertionError(name)


def test_change_flow_with_upgrade_undo_and_refresh(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    from dpg.core.actions.model import ActionKind
    from dpg.core.audit.model import Exposure
    from dpg.core.auth.scopes import AccessLevel as Level
    from dpg.gui.actions_ui import ActionDialog, ResultDialog

    assert env.ctx is not None
    _login(env)
    _writable(env)
    ids = _populate(env.fake)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    _audit(qtbot, window)

    upgrades: list[Level] = []

    def fake_setup(level: Level = Level.AUDIT) -> bool:
        upgrades.append(level)
        env.ctx.manager.login(level)  # type: ignore[union-attr]
        return True

    monkeypatch.setattr(window, "run_setup", fake_setup)
    shown: list[Any] = []

    def show(dialog: Any) -> int:
        shown.append(dialog)
        if isinstance(dialog, ActionDialog):
            assert dialog.kind.currentData() == ActionKind.REMOVE_LINK
            assert len(dialog.plan.changes) == 1
            return 1
        return 0

    env.ctx.show_dialog = show
    _select_file(window, "공지.pdf")
    window.change_selected()
    qtbot.waitUntil(
        lambda: window.task is None and any(isinstance(d, ResultDialog) for d in shown),
        timeout=15000,
    )
    qtbot.waitUntil(lambda: window.task is None, timeout=15000)  # automatic re-audit
    assert upgrades == [Level.MODIFY]
    assert env.fake.visibility(env.fake.items[ids["link"]]) == "limited"
    assert env.fake.notifications == []
    link_row = next(a for a in window.model.items if a.file_id == ids["link"])
    assert link_row.exposure == Exposure.RESTRICTED  # table shows the new state

    # undo from the history dialog
    def show_history(dialog: Any) -> int:
        shown.append(dialog)
        if type(dialog).__name__ == "HistoryDialog":
            dialog.table.selectRow(0)
            dialog._undo()
        return 1

    env.ctx.show_dialog = show_history
    window.open_history()
    qtbot.waitUntil(
        lambda: (
            window.task is None
            and env.fake.visibility(env.fake.items[ids["link"]]) == "anyoneWithLink"
        ),
        timeout=15000,
    )
    qtbot.waitUntil(lambda: window.task is None, timeout=15000)
    results = [d for d in shown if isinstance(d, ResultDialog)]
    assert "되돌림" in results[-1].findChild(QTableWidget).item(0, 2).text()


def test_dry_run_changes_nothing(qtbot: Any, env: Env) -> None:
    from dpg.core.auth.scopes import AccessLevel as Level
    from dpg.gui.actions_ui import ActionDialog

    assert env.ctx is not None
    _login(env)
    env.ctx.manager.login(Level.MODIFY)
    _writable(env)
    ids = _populate(env.fake)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    _audit(qtbot, window)

    def show(dialog: Any) -> int:
        if isinstance(dialog, ActionDialog):
            dialog.dry_run.setChecked(True)
            return 1
        return 0

    env.ctx.show_dialog = show
    _select_file(window, "공지.pdf")
    window.change_selected()
    qtbot.waitUntil(
        lambda: window.task is None and "드라이런" in window.progress_label.text(), timeout=15000
    )
    assert env.fake.write_calls == []
    assert env.fake.visibility(env.fake.items[ids["link"]]) == "anyoneWithLink"


def test_many_files_need_confirmation_phrase(qtbot: Any) -> None:
    from PySide6.QtWidgets import QDialogButtonBox

    from dpg.core.actions.model import ActionKind, Change, Op, Plan
    from dpg.core.actions.planner import CONFIRM_PHRASE
    from dpg.gui.actions_ui import ActionDialog

    def build(kind: ActionKind) -> Plan:
        plan = Plan(kind)
        for i in range(51):
            plan.changes.append(
                Change(
                    f"f{i}",
                    Op.DELETE_PERM,
                    "anyoneWithLink",
                    {"type": "anyone", "role": "reader"},
                    None,
                    "링크 제거",
                )
            )
        return plan

    dialog = ActionDialog(build, {}, ActionKind.REMOVE_LINK, False)
    qtbot.addWidget(dialog)
    ok = dialog.buttons.button(QDialogButtonBox.StandardButton.Ok)
    assert not ok.isEnabled()
    dialog.confirm.setText("권한 변경")
    assert not ok.isEnabled()
    dialog.confirm.setText(CONFIRM_PHRASE)
    assert ok.isEnabled()
    dialog.dry_run.setChecked(True)  # a dry run changes nothing: no phrase needed
    dialog.confirm.setText("")
    assert ok.isEnabled()


def test_change_requires_consent_for_extra_permission(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    _login(env)
    _populate(env.fake)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    _audit(qtbot, window)
    env.ctx.confirm = lambda *a: False
    _select_file(window, "공지.pdf")
    window.change_selected()
    assert window.task is None
    assert env.fake.write_calls == []


def test_switching_action_in_dropdown_builds_the_right_plan(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Regression: QComboBox returns the plain string value, not the ActionKind member.
    Every action must still plan correctly after the user changes the dropdown."""
    from dpg.core.actions.model import ActionKind
    from dpg.core.auth.scopes import AccessLevel as Level
    from dpg.gui.actions_ui import ActionDialog

    assert env.ctx is not None
    _login(env)
    env.ctx.manager.login(Level.MODIFY)
    _writable(env)
    _populate(env.fake)
    folder = env.fake.add_folder("f-하위폴더")
    env.fake.share(folder, "user", "reader", email="friend@gmail.example")
    child = env.fake.add_file("e-상속", parent=folder)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    _audit(qtbot, window)
    captured: list[ActionDialog] = []

    def show(dialog: Any) -> int:
        if isinstance(dialog, ActionDialog):
            captured.append(dialog)
            dialog.kind.setCurrentIndex(dialog.kind.findData(ActionKind.REMOVE_EXTERNAL))
            return 1
        return 0

    env.ctx.show_dialog = show
    # inherited share on the child: switched to "remove external" -> excluded with the right reason
    _select_file(window, "e-상속")
    window.change_selected()
    dialog = captured[-1]
    assert dialog.plan.action is ActionKind.REMOVE_EXTERNAL
    assert dialog.plan.changes == []
    assert "상위 폴더" in dialog.plan.skipped[0].reason
    qtbot.waitUntil(lambda: window.task is None, timeout=15000)

    # every action works after switching (plan.action is always the enum member)
    from dpg.core.actions.model import PERMISSION_ACTIONS

    for kind in PERMISSION_ACTIONS:
        dialog.kind.setCurrentIndex(dialog.kind.findData(kind))
        assert dialog.plan.action is kind

    # on the folder: switched action executes and is recorded
    _select_file(window, "📁 f-하위폴더 /")
    window.change_selected()
    qtbot.waitUntil(
        lambda: (
            window.task is None
            and not any(
                p.email == "friend@gmail.example" for p in env.fake.items[folder].permissions
            )
        ),
        timeout=15000,
    )
    qtbot.waitUntil(lambda: window.task is None, timeout=15000)
    assert not any(
        p["type"] == "user" and p.get("emailAddress") == "friend@gmail.example"
        for p in env.fake.effective_permissions(env.fake.items[child])
    )


# --- big drives -------------------------------------------------------------------------------


def test_live_refresh_can_be_switched_off_on_screen(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    _login(env)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    assert env.ctx.prefs.live_refresh
    window.live_btn.click()
    assert not env.ctx.prefs.live_refresh
    assert not window.live_timer.isActive()
    assert "꺼짐" in window.live_label.text()
    assert window.live_btn.text() == "켜기"
    window.live_btn.click()
    assert env.ctx.prefs.live_refresh


def test_live_refresh_pauses_on_a_big_drive(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert env.ctx is not None
    _login(env)
    _populate(env.fake)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    _audit(qtbot, window)
    assert window.result is not None
    monkeypatch.setattr(mw, "LIVE_MAX_ITEMS", 2)
    env.fake.add_file("새 파일.txt")
    window._live_tick()
    assert window._live_check is None  # did not even ask Drive
    assert "자동 반영 쉼" in window.live_label.text()


def test_detection_of_shared_files_only(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert env.ctx is not None
    _login(env)
    env.ctx.manager.login(AccessLevel.DETECT)
    ids = _populate(env.fake)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    window.detect_check.setChecked(True)
    assert window.detect_shared_check.isEnabled()
    window.detect_shared_check.setChecked(True)
    _audit(qtbot, window)
    assert window.detections is not None
    assert ids["link"] in window.detections
    assert ids["ext"] in window.detections
    assert ids["private"] not in window.detections


def test_scan_keeps_the_computer_awake_and_shows_progress(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert env.ctx is not None
    _login(env)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    calls: list[str] = []
    monkeypatch.setattr(window.keep_awake, "start", lambda: calls.append("start"))
    monkeypatch.setattr(window.keep_awake, "stop", lambda: calls.append("stop"))
    _audit(qtbot, window)
    assert calls[:2] == ["start", "stop"]
    window._on_progress(mw.OfflineWait(125))
    assert "인터넷 연결이 끊겼습니다" in window.progress_label.text()
    assert "2분" in window.progress_label.text()
    window._on_progress(
        mw.Progress("list", listed=12000, window=3, windows=31, window_label="2013년 상반기")
    )
    assert "구간 3/31 (2013년 상반기)" in window.progress_label.text()


def test_public_scope_is_offered(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    _login(env)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    window.scope_combo.setCurrentIndex(window.scope_combo.findData("public"))
    scope = window.current_scope()
    assert scope is not None
    assert scope.kind == "public"


def test_duration_wording() -> None:
    assert mw._duration_ko(30) == "1분"
    assert mw._duration_ko(59 * 60) == "59분"
    assert mw._duration_ko(3 * 3600 + 20 * 60) == "3시간 20분"
    assert mw._duration_ko(3 * 86400 + 3600) == "3일 1시간"


# --- scheduled scans (D-099) ------------------------------------------------------------------


def _scheduled_window(env: Env, moment: dt.datetime, **kw: Any) -> MainWindow:
    from dpg.core.schedule import Schedule

    assert env.ctx is not None
    env.ctx.prefs.schedule = Schedule(enabled=True, start="22:00", end="06:00", **kw).to_json()
    window = MainWindow(env.ctx)
    window.now = lambda: moment
    return window


def test_schedule_dialog_saves_and_turns_on_the_login_item(qtbot: Any, env: Env) -> None:
    from dpg.gui.schedule_ui import ScheduleDialog

    assert env.ctx is not None
    dialog = ScheduleDialog(env.ctx, now=dt.datetime(2026, 10, 5, 9, 0))
    qtbot.addWidget(dialog)
    dialog.enabled.setChecked(True)
    for i, box in enumerate(dialog.days):
        box.setChecked(i < 5)
    dialog.detect.setChecked(True)
    dialog.autostart.setChecked(True)
    assert "다음 예약: 오늘 22:00" in dialog.preview.text()
    dialog._save()
    saved = env.ctx.prefs.schedule
    assert saved["enabled"]
    assert saved["days"] == [0, 1, 2, 3, 4]
    assert saved["detect"]
    assert env.ctx.autostart.on


def test_schedule_dialog_reports_a_login_item_it_cannot_make(qtbot: Any, env: Env) -> None:
    from dpg.gui.autostart import AutostartError
    from dpg.gui.schedule_ui import ScheduleDialog

    assert env.ctx is not None
    env.ctx.autostart.fail = AutostartError("앱이 임시 위치에서 실행 중입니다.")
    dialog = ScheduleDialog(env.ctx)
    qtbot.addWidget(dialog)
    dialog.enabled.setChecked(True)
    dialog.autostart.setChecked(True)
    dialog._save()
    assert env.notices[-1] == ("자동으로 켜기", "앱이 임시 위치에서 실행 중입니다.")
    assert not env.ctx.prefs.schedule.get("enabled")  # nothing saved: the user can fix it


def test_scheduled_scan_runs_once_per_window(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    _login(env)
    _populate(env.fake)
    window = _scheduled_window(env, dt.datetime(2026, 10, 5, 23, 0))
    qtbot.addWidget(window)
    window._schedule_tick()
    assert window.task is not None
    assert window._sched_run
    qtbot.waitUntil(lambda: window.task is None, timeout=15000)
    assert window._sched_done_for == dt.datetime(2026, 10, 5, 22, 0)
    assert "예약 검사 완료" in window.progress_label.text()
    assert "링크 공개 2" in window.progress_label.text()
    window._schedule_tick()
    assert window.task is None  # already done for tonight
    window.now = lambda: dt.datetime(2026, 10, 6, 22, 30)  # the next night
    window._schedule_tick()
    assert window.task is not None
    qtbot.waitUntil(lambda: window.task is None, timeout=15000)


def test_scheduled_scan_stops_when_its_window_ends(qtbot: Any, env: Env) -> None:
    from dpg.core.audit.runner import AuditCancelled

    assert env.ctx is not None
    _login(env)
    window = _scheduled_window(env, dt.datetime(2026, 10, 6, 6, 1))
    qtbot.addWidget(window)
    window.task = object()  # type: ignore[assignment]  # a scheduled scan still running
    window._sched_run = True
    window._schedule_tick()
    assert window.cancel_event.is_set()
    window._on_failed(AuditCancelled())
    assert "다음 예약 때 그 자리부터 이어서" in window.progress_label.text()
    window.task = None


def test_scheduled_scan_without_login_only_tells(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    window = _scheduled_window(env, dt.datetime(2026, 10, 5, 23, 0))
    qtbot.addWidget(window)
    window._schedule_tick()
    assert window.task is None
    assert "로그인되어 있지 않아" in window.progress_label.text()
    assert env.notices == []  # no dialog popping up in the middle of the night


def test_scheduled_detection_without_permission_is_skipped_and_said(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    _login(env)  # audit permission only
    _populate(env.fake)
    window = _scheduled_window(env, dt.datetime(2026, 10, 5, 23, 0), detect=True)
    qtbot.addWidget(window)
    window._schedule_tick()
    qtbot.waitUntil(lambda: window.task is None, timeout=15000)
    assert window.detections is None
    assert "개인정보는 권한이 없어 건너뜀" in window.progress_label.text()


def test_tray_keeps_the_app_running_when_the_window_closes(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The menu-bar icon path (offscreen has no tray, so pretend there is one)."""
    from PySide6.QtGui import QCloseEvent
    from PySide6.QtWidgets import QSystemTrayIcon

    monkeypatch.setattr(QSystemTrayIcon, "isSystemTrayAvailable", staticmethod(lambda: True))
    assert env.ctx is not None
    _login(env)
    window = _scheduled_window(env, dt.datetime(2026, 10, 5, 9, 0))
    qtbot.addWidget(window)
    assert window.tray is not None
    assert window.tray.isVisible()
    assert "매일 22:00~06:00" in window.tray.toolTip()
    assert not window.tray.icon().isNull()
    window.show()
    event = QCloseEvent()
    window.closeEvent(event)
    assert not event.isAccepted()  # hidden, still running for the schedule
    assert not window.isVisible()
    window.show_window()
    assert window.isVisible()
    window._quitting = True
    event = QCloseEvent()
    window.closeEvent(event)
    assert event.isAccepted()
    assert not window.tray.isVisible()


def test_no_schedule_means_no_tray_and_a_normal_close(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    from PySide6.QtGui import QCloseEvent
    from PySide6.QtWidgets import QSystemTrayIcon

    monkeypatch.setattr(QSystemTrayIcon, "isSystemTrayAvailable", staticmethod(lambda: True))
    assert env.ctx is not None
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    assert window.tray is not None
    assert not window.tray.isVisible()
    event = QCloseEvent()
    window.closeEvent(event)
    assert event.isAccepted()


def test_personal_data_check_of_chosen_files_only(
    qtbot: Any, env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    from dpg.core.auth.scopes import AccessLevel as Level
    from dpg.core.policy import DetectStatus

    assert env.ctx is not None
    _login(env)
    fixtures = Path(__file__).resolve().parents[1] / "fixtures" / "synthetic"
    csv_id = env.fake.add_file(
        "6-2_학생_연락처.csv",
        mime_type="text/csv",
        content=(fixtures / "6-2_학생_연락처.csv").read_bytes(),
    )
    env.fake.share(csv_id, "anyone", "reader")
    other = env.fake.add_file("메모.txt", content=b"hello")
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    _audit(qtbot, window)  # sharing only, no detection
    assert window.detections is None
    upgrades: list[Level] = []

    def fake_setup(level: Level = Level.AUDIT) -> bool:
        upgrades.append(level)
        env.ctx.manager.login(level)  # type: ignore[union-attr]
        return True

    monkeypatch.setattr(window, "run_setup", fake_setup)
    window.set_filter("link")
    window.model.checked = {csv_id}
    window._update_actions()
    assert window.detect_btn.isEnabled()
    gets = env.fake.call_count("files.get_media")
    window.detect_btn.click()
    qtbot.waitUntil(lambda: window.task is None, timeout=15000)
    assert upgrades == [Level.DETECT]
    assert window.detections is not None
    assert window.detections[csv_id].status is DetectStatus.DETECTED
    assert other not in window.detections  # only what was chosen
    assert env.fake.call_count("files.get_media") == gets + 1
    assert "개인정보 있음 1개" in window.progress_label.text()
    window.detect_btn.click()  # again: nothing is downloaded twice
    qtbot.waitUntil(lambda: window.task is None, timeout=15000)
    assert env.fake.call_count("files.get_media") == gets + 1


def test_open_chosen_files_in_drive(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    _login(env)
    ids = _populate(env.fake)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    _audit(qtbot, window)
    window.model.checked = {ids["link"], ids["ext"]}
    window._update_actions()
    assert window.open_btn.isEnabled()
    window.open_btn.click()
    assert f"https://drive.google.com/open?id={ids['link']}".split("?")[0] in env.opened
    assert env.opened.count("https://drive.google.com/open") == 2


def test_inherited_link_offers_the_parent_folder(qtbot: Any, env: Env) -> None:
    from dpg.core.actions.model import ActionKind
    from dpg.core.actions.planner import build_plan
    from dpg.gui.actions_ui import ActionDialog

    assert env.ctx is not None
    _login(env)
    folder = env.fake.add_folder("공개 폴더")
    env.fake.share(folder, "anyone", "reader")
    child = env.fake.add_file("계약서.docx", parent=folder)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    _audit(qtbot, window)
    assert window.result is not None
    items = [a for a in window.result.items if a.file_id == child]
    names = {a.file_id: a.name for a in window.result.items}

    def build(kind: ActionKind) -> Any:
        return build_plan(
            kind, items, account=ME, internal_domains=(), all_items=window.result.items
        )  # type: ignore[union-attr]

    dialog = ActionDialog(build, names, ActionKind.REMOVE_LINK, False)
    qtbot.addWidget(dialog)
    assert not dialog.plan.changes
    assert dialog.plan.skipped[0].folder_id == folder
    assert "공개 폴더" in dialog.folder_btn.text()
    dialog.folder_btn.click()
    assert dialog.go_to_folders == [folder]


def test_change_dialog_moves_on_to_the_parent_folder(qtbot: Any, env: Env) -> None:
    assert env.ctx is not None
    _login(env)
    env.ctx.manager.login(AccessLevel.MODIFY)
    folder = env.fake.add_folder("공개 폴더")
    env.fake.share(folder, "anyone", "reader")
    child = env.fake.add_file("계약서.docx", parent=folder)
    window = MainWindow(env.ctx)
    qtbot.addWidget(window)
    _audit(qtbot, window)
    shown: list[list[str]] = []

    def show(dialog: Any) -> bool:
        shown.append([s.file_id for s in dialog.plan.skipped] + dialog.plan.file_ids)
        if len(shown) == 1:
            dialog.folder_btn.click()  # "상위 폴더에서 바꾸기"
        return False

    env.ctx.show_dialog = show
    window.model.checked = {child}
    window.change_selected()
    assert len(shown) == 2
    assert shown[1] == [folder]  # the second preview is for the folder itself
