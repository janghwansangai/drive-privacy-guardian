"""Synthetic personal-data fixture generator (SPEC 8).

Every value is fake *by construction* so it can never belong to a real person
(rules recorded in DECISIONS.md D-012; enforced by `find_non_synthetic`):

- 주민등록번호: 7th digit 3/4 with birth year 2050–2099 (a future date → no living person).
  Checksum-valid unless deliberately generated as a "post-2020 style" (checksum-invalid) value.
- 휴대전화: 010-0xxx-xxxx / 010-1xxx-xxxx (middle numbers starting with 0/1 are not assigned).
- 일반전화: local number starting with 0/1 (never assigned).
- 이메일: RFC 2606 reserved domains only (example.com/.org/.net).
- 카드번호: 9999 prefix + Luhn check digit.
- 계좌번호: 999 prefix. 여권번호: M0000xxxx / M000A0000 pattern.
- 주소: real 시·도 + non-existent 구 and road names (가상구, 샘플로, ...).
- 이름: random syllable combinations (not linked to any real identifier).

Usage:
    uv run python tools/gen_synthetic.py --out tests/fixtures/synthetic --seed 20260926
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import random
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from dpg.core.detect.validators import luhn_check_digit, rrn_birth_date, rrn_checksum

if __package__ in (None, ""):  # run as a script: make `tools.*` importable
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

GENERATOR_VERSION = 2
FUTURE_YEAR_MIN = 2050

SURNAMES = "김이박최정강조윤장임한오서신권황안송류홍"
GIVEN = "가나다라마바사아자차카타파하온솔빛결담별윤린준율"
SIDO = [
    "서울특별시",
    "부산광역시",
    "대구광역시",
    "인천광역시",
    "광주광역시",
    "대전광역시",
    "경기도",
    "강원특별자치도",
    "충청북도",
    "전라남도",
    "경상북도",
    "제주특별자치도",
]
FAKE_GU = ["가상구", "샘플구", "예시구", "테스트구", "허구구"]
FAKE_ROAD = ["가상로", "샘플로", "예시대로", "테스트길", "허구로"]
BANKS = ["국민은행", "신한은행", "우리은행", "하나은행", "농협은행"]
RESERVED_EMAIL_DOMAINS = ("example.com", "example.org", "example.net")
SENSITIVE_KEYWORDS = ["상담", "진단", "복약", "학교폭력"]


# ---------------------------------------------------------------------------------------------
# Value generators
# ---------------------------------------------------------------------------------------------


def fake_name(rng: random.Random) -> str:
    return rng.choice(SURNAMES) + rng.choice(GIVEN) + rng.choice(GIVEN)


def fake_birth(rng: random.Random) -> tuple[int, int, int]:
    return rng.randint(FUTURE_YEAR_MIN, 2099), rng.randint(1, 12), rng.randint(1, 28)


def fake_rrn(
    rng: random.Random, *, checksum_valid: bool = True, birth: tuple[int, int, int] | None = None
) -> str:
    y, m, d = birth or fake_birth(rng)
    gender = rng.choice("34")
    first12 = f"{y % 100:02d}{m:02d}{d:02d}{gender}{rng.randint(0, 99999):05d}"
    check = rrn_checksum(first12)
    if not checksum_valid:
        check = (check + rng.randint(1, 9)) % 10
    return f"{first12[:6]}-{first12[6:]}{check}"


def fake_mobile(rng: random.Random) -> str:
    return f"010-{rng.randint(0, 1999):04d}-{rng.randint(0, 9999):04d}"


def fake_landline(rng: random.Random) -> str:
    area = rng.choice(["02", "031", "051", "062"])
    return f"{area}-{rng.randint(0, 1999):04d}-{rng.randint(0, 9999):04d}"


def fake_email(rng: random.Random) -> str:
    user = "".join(rng.choice("abcdefghijklmnopqrstuvwxyz") for _ in range(7))
    return f"{user}@{rng.choice(RESERVED_EMAIL_DOMAINS)}"


def fake_card(rng: random.Random) -> str:
    payload = "9999" + "".join(str(rng.randint(0, 9)) for _ in range(11))
    full = payload + str(luhn_check_digit(payload))
    return "-".join(full[i : i + 4] for i in range(0, 16, 4))


def fake_account(rng: random.Random) -> str:
    return f"999-{rng.randint(10, 99)}-{rng.randint(0, 999999):06d}"


def fake_passport(rng: random.Random) -> str:
    if rng.random() < 0.5:
        return f"M0000{rng.randint(0, 9999):04d}"
    return f"M000{rng.choice('ABCDEFGH')}{rng.randint(0, 9999):04d}"


def fake_license(rng: random.Random) -> str:
    region, year = rng.randint(11, 28), rng.randint(0, 99)
    return f"{region}-{year:02d}-{rng.randint(0, 999999):06d}-{rng.randint(0, 99):02d}"


def fake_address(rng: random.Random) -> str:
    return f"{rng.choice(SIDO)} {rng.choice(FAKE_GU)} {rng.choice(FAKE_ROAD)} {rng.randint(1, 300)}"


# ---------------------------------------------------------------------------------------------
# Fixture set
# ---------------------------------------------------------------------------------------------


@dataclass
class Manifest:
    seed: int
    generator_version: int = GENERATOR_VERSION
    # file name -> kind -> expected count (ground truth for Phase 4 precision/recall)
    files: dict[str, dict[str, int]] = field(default_factory=dict)
    # file name -> expected "검사 불가" reason (must never be reported as clean)
    unscannable: dict[str, str] = field(default_factory=dict)

    def to_json(self) -> str:
        return (
            json.dumps(
                {
                    "seed": self.seed,
                    "generator_version": self.generator_version,
                    "files": {k: dict(sorted(v.items())) for k, v in sorted(self.files.items())},
                    "unscannable": dict(sorted(self.unscannable.items())),
                },
                ensure_ascii=False,
                indent=2,
            )
            + "\n"
        )


ROSTER_HEADER = ["번호", "성명", "생년월일", "주민등록번호", "보호자 연락처", "주소"]


def _roster_rows(rng: random.Random, n: int) -> list[list[str]]:
    rows = []
    for i in range(1, n + 1):
        birth = fake_birth(rng)
        rows.append(
            [
                str(i),
                fake_name(rng),
                f"{birth[0]}-{birth[1]:02d}-{birth[2]:02d}",
                fake_rrn(rng, birth=birth),
                fake_mobile(rng),
                fake_address(rng),
            ]
        )
    return rows


def _notice_lines(rng: random.Random) -> tuple[list[str], Counter[str]]:
    c: Counter[str] = Counter()
    lines = ["가정통신문 — 현장체험학습 안내 (합성 데이터)"]
    lines.append(f"담임 연락처: {fake_mobile(rng)}, 교무실 {fake_landline(rng)}")
    c["mobile"] += 1
    c["landline"] += 1
    lines.append(f"문의 메일: {fake_email(rng)}")
    c["email"] += 1
    for _ in range(3):
        lines.append(
            f"환불 계좌: {rng.choice(BANKS)} {fake_account(rng)} (예금주 {fake_name(rng)})"
        )
        c["account"] += 1
    lines.append("참가비는 2026-10-15까지 납부해 주세요. 주문번호 20261015-0001.")
    return lines, c


def _mixed_lines(rng: random.Random) -> tuple[list[str], Counter[str]]:
    c: Counter[str] = Counter()
    lines = ["업무 메모 (합성 데이터)"]
    for _ in range(2):
        lines.append(f"여권번호 {fake_passport(rng)} 사본 제출 완료")
        c["passport"] += 1
    lines.append(f"운전면허 번호: {fake_license(rng)}")
    c["driver_license"] += 1
    lines.append(f"법인카드 {fake_card(rng)} 사용 내역 정리")
    c["card"] += 1
    for _ in range(2):
        lines.append(f"주민번호 {fake_rrn(rng)} 확인")
        c["rrn"] += 1
    lines.append(f"신규 발급 번호(검증식 불일치) {fake_rrn(rng, checksum_valid=False)}")
    c["rrn_suspect"] += 1
    lines.append(f"{fake_name(rng)} 학생 {rng.choice(SENSITIVE_KEYWORDS)} 기록 — 담당자 확인 필요")
    c["sensitive_suspect"] += 1
    lines.append(f"자택 주소: {fake_address(rng)}")
    c["address"] += 1
    return lines, c


NEGATIVE_LINES = [
    "비개인정보 샘플 (정밀도 측정용)",
    "회의 일시: 2026-09-26 14:00",
    "대표번호 1588-1234 로 문의",
    "주문번호 20260926-0001, 20260926-0002",
    "재고 코드 12345678901 (문맥 없는 긴 숫자열)",
    "학번 2026123456, 출석 번호 17",
    "날짜가 아닌 13자리: 991332-1234567",
    "버전 3.12.14, 우편번호 04524",
    "학생 수 30명, 평균 점수 87.5",
]


def generate(out_dir: Path, seed: int, roster_size: int = 30) -> Manifest:
    rng = random.Random(seed)
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest = Manifest(seed=seed)

    rows = _roster_rows(rng, roster_size)
    roster_counts = {
        "rrn": roster_size,
        "mobile": roster_size,
        "address": roster_size,
        "student_roster": 1,
    }

    csv_name = "6-2_학생_연락처.csv"
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(ROSTER_HEADER)
    writer.writerows(rows)
    (out_dir / csv_name).write_text(buf.getvalue(), encoding="utf-8")
    manifest.files[csv_name] = dict(roster_counts)

    xlsx_name = "6-2_학생_연락처.xlsx"
    _write_xlsx(out_dir / xlsx_name, [ROSTER_HEADER, *rows])
    manifest.files[xlsx_name] = dict(roster_counts)

    notice, notice_counts = _notice_lines(rng)
    docx_name = "가정통신문_체험학습.docx"
    _write_docx(out_dir / docx_name, notice)
    manifest.files[docx_name] = dict(notice_counts)

    mixed, mixed_counts = _mixed_lines(rng)
    txt_name = "업무메모.txt"
    (out_dir / txt_name).write_text("\n".join(mixed) + "\n", encoding="utf-8")
    manifest.files[txt_name] = dict(mixed_counts)

    _generate_v2(rng, out_dir, manifest)

    neg_name = "negatives.txt"
    (out_dir / neg_name).write_text("\n".join(NEGATIVE_LINES) + "\n", encoding="utf-8")
    manifest.files[neg_name] = {}

    (out_dir / "manifest.json").write_text(manifest.to_json(), encoding="utf-8")
    return manifest


def _generate_v2(rng: random.Random, out_dir: Path, manifest: Manifest) -> None:
    """HWP / HWPX / PDF fixtures and files that must come out as 검사 불가."""
    from tools.doc_writers import build_hwp, build_hwpx, build_pdf

    # HWP: counselling note with a contact table
    table = [["번호", "성명", "연락처"]]
    table += [[str(i), fake_name(rng), fake_mobile(rng)] for i in range(1, 7)]
    hwp_blocks: list[str | list[list[str]]] = [
        "학생 상담 기록 (합성 데이터)",
        f"{fake_name(rng)} 학생 상담 내용: 교우 관계 고민",
        f"보호자 연락처 {fake_mobile(rng)}",
        f"주민번호 {fake_rrn(rng)}",
        table,
    ]
    (out_dir / "상담기록_가상.hwp").write_bytes(build_hwp(hwp_blocks))
    manifest.files["상담기록_가상.hwp"] = {
        "sensitive_suspect": 1,
        "mobile": 7,
        "rrn": 1,
        "student_roster": 1,
    }

    # HWPX: class roster with RRN and address
    roster = [["번호", "성명", "생년월일", "주민등록번호", "주소"]]
    for i in range(1, 9):
        birth = fake_birth(rng)
        roster.append(
            [
                str(i),
                fake_name(rng),
                f"{birth[0]}-{birth[1]:02d}-{birth[2]:02d}",
                fake_rrn(rng, birth=birth),
                fake_address(rng),
            ]
        )
    (out_dir / "학생명단_가상.hwpx").write_bytes(
        build_hwpx(["6학년 2반 명단 (합성 데이터)", f"담당 {fake_email(rng)}", roster])
    )
    manifest.files["학생명단_가상.hwpx"] = {
        "rrn": 8,
        "address": 8,
        "student_roster": 1,
        "email": 1,
    }

    # PDF (text)
    lines = ["보호자 안내문 (합성 데이터)"]
    lines += [f"환불 계좌: {rng.choice(BANKS)} {fake_account(rng)}" for _ in range(2)]
    lines += [f"문의 {fake_mobile(rng)}", f"주소: {fake_address(rng)}"]
    (out_dir / "보호자_안내문.pdf").write_bytes(build_pdf([lines]))
    manifest.files["보호자_안내문.pdf"] = {"account": 2, "mobile": 1, "address": 1}

    # Must be 검사 불가, never "clean"
    (out_dir / "스캔_가상.pdf").write_bytes(build_pdf([[], []]))
    manifest.unscannable["스캔_가상.pdf"] = "image_only"
    (out_dir / "암호_가상.hwp").write_bytes(build_hwp([f"주민번호 {fake_rrn(rng)}"], password=True))
    manifest.unscannable["암호_가상.hwp"] = "encrypted"
    (out_dir / "배포용_가상.hwp").write_bytes(
        build_hwp([f"주민번호 {fake_rrn(rng)}"], distribution=True)
    )
    manifest.unscannable["배포용_가상.hwp"] = "distribution"


def _write_xlsx(path: Path, rows: list[list[str]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "명단"
    for row in rows:
        ws.append(row)
    wb.save(path)


def _write_docx(path: Path, lines: list[str]) -> None:
    from docx import Document

    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(str(path))


# ---------------------------------------------------------------------------------------------
# Fakeness checker — proves fixtures contain no potentially real identifiers
# ---------------------------------------------------------------------------------------------

_RRN_SHAPE = re.compile(r"(?<!\d)(\d{6})-?(\d{7})(?!\d)")
_MOBILE = re.compile(r"(?<!\d)01[016789]-?(\d{3,4})-?(\d{4})(?!\d)")
_LANDLINE = re.compile(r"(?<!\d)0(?:2|[3-6][1-5])-(\d{3,4})-(\d{4})(?!\d)")
_EMAIL = re.compile(r"[A-Za-z0-9._%+\-]+@([A-Za-z0-9.\-]+\.[A-Za-z]{2,})")
_CARD = re.compile(r"(?<!\d)(\d{4})-\d{4}-\d{4}-\d{4}(?!\d)")


def find_non_synthetic(text: str) -> list[str]:
    """Return reasons (never values) for any identifier that could belong to a real person."""
    problems: list[str] = []
    for m in _RRN_SHAPE.finditer(text):
        born = rrn_birth_date(m.group(1) + m.group(2))
        if born is not None and born.year < FUTURE_YEAR_MIN:
            problems.append(f"rrn-shaped value with plausible birth year at offset {m.start()}")
    for m in _MOBILE.finditer(text):
        if m.group(1)[0] not in "01":
            problems.append(f"mobile number in assignable range at offset {m.start()}")
    for m in _LANDLINE.finditer(text):
        if m.group(1)[0] not in "01":
            problems.append(f"landline number in assignable range at offset {m.start()}")
    for m in _EMAIL.finditer(text):
        if m.group(1).lower() not in RESERVED_EMAIL_DOMAINS:
            problems.append(f"email with non-reserved domain at offset {m.start()}")
    for m in _CARD.finditer(text):
        if m.group(1) != "9999":
            problems.append(f"card number outside 9999 prefix at offset {m.start()}")
    return problems


def fixture_text(path: Path) -> str:
    """Plain text of a generated fixture (CSV/TXT/XLSX/DOCX)."""
    suffix = path.suffix.lower()
    if suffix in (".csv", ".txt", ".json"):
        return path.read_text(encoding="utf-8")
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True)
        return "\n".join(
            "\t".join("" if v is None else str(v) for v in row)
            for ws in wb.worksheets
            for row in ws.iter_rows(values_only=True)
        )
    if suffix == ".docx":
        from docx import Document

        return "\n".join(p.text for p in Document(str(path)).paragraphs)
    if suffix in (".hwp", ".hwpx", ".pdf"):
        from dpg.core.extract import Unscannable, extract, plan_for

        try:
            doc = extract(path.read_bytes(), plan_for("", path.name))
        except Unscannable:
            return _raw_text(path.read_bytes())
        return "\n".join(
            [s.text for s in doc.segments] + ["\t".join(r) for t in doc.tables for r in t.rows]
        )
    raise ValueError(f"unsupported fixture type: {suffix}")


def _raw_text(data: bytes) -> str:
    """Files the extractor refuses (encrypted/scanned) are still checked for fake-ness by
    decoding their raw (and, where possible, decompressed) bytes."""
    import zlib

    chunks = [data]
    for start in range(0, len(data), 512):
        try:
            chunks.append(zlib.decompressobj(-15).decompress(data[start:]))
        except zlib.error:
            continue
    return "\n".join(c.decode("utf-16-le", "ignore") + c.decode("utf-8", "ignore") for c in chunks)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="합성 개인정보 테스트 데이터 생성기")
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--seed", type=int, default=20260926)
    parser.add_argument("--roster-size", type=int, default=30)
    args = parser.parse_args(argv)
    manifest = generate(args.out, args.seed, args.roster_size)
    bad = [
        (name, p)
        for name in [*manifest.files, *manifest.unscannable]
        for p in find_non_synthetic(fixture_text(args.out / name))
    ]
    for name, problem in bad:
        print(f"NOT SYNTHETIC: {name}: {problem}", file=sys.stderr)
    print(f"generated {len(manifest.files)} files in {args.out}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
