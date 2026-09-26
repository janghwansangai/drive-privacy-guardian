"""Extractors for TXT, CSV, DOCX, XLSX, PDF, HWPX. (HWP 5.0 lives in hwp.py.)"""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from dpg.core.extract.base import (
    MAX_TABLE_ROWS,
    MAX_TEXT_CHARS,
    Extracted,
    Table,
    Unscannable,
    UnscannableReason,
    check_size,
    decode_text,
    open_zip_safely,
)

# --- plain text ---------------------------------------------------------------------------------


def extract_txt(data: bytes) -> Extracted:
    check_size(data)
    out = Extracted()
    text = decode_text(data)[:MAX_TEXT_CHARS]
    for i, line in enumerate(text.splitlines(), start=1):
        out.add_text(line, f"{i}행")
    return out


def extract_csv(data: bytes, delimiter: str | None = None) -> Extracted:
    check_size(data)
    text = decode_text(data)[:MAX_TEXT_CHARS]
    if delimiter is None:
        sample = text[:4096]
        delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    rows: list[list[str]] = []
    out = Extracted()
    try:
        for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
            if i > MAX_TABLE_ROWS:
                out.truncated = True
                break
            rows.append([c.strip() for c in row])
    except csv.Error:
        raise Unscannable(UnscannableReason.CORRUPT) from None
    out.tables.append(Table("표", rows))
    return out


# --- office ------------------------------------------------------------------------------------


def extract_docx(data: bytes) -> Extracted:
    check_size(data)
    open_zip_safely(data).close()
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception:  # noqa: BLE001 — library raises many types; never echo content
        raise Unscannable(UnscannableReason.CORRUPT) from None
    out = Extracted()
    for i, p in enumerate(doc.paragraphs, start=1):
        out.add_text(p.text, f"문단 {i}")
    for t_index, table in enumerate(doc.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows[:MAX_TABLE_ROWS]]
        out.tables.append(Table(f"표 {t_index}", rows))
    return out


def extract_xlsx(data: bytes) -> Extracted:
    check_size(data)
    open_zip_safely(data).close()
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise Unscannable(UnscannableReason.CORRUPT) from None
    out = Extracted()
    try:
        for s_index, ws in enumerate(wb.worksheets, start=1):
            rows: list[list[str]] = []
            for r_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if r_index > MAX_TABLE_ROWS:
                    out.truncated = True
                    break
                rows.append(["" if v is None else str(v).strip() for v in row])
            # Location uses the sheet number, not its name (a name could identify a person).
            out.tables.append(Table(f"시트 {s_index}", rows))
    finally:
        wb.close()
    return out


# --- PDF ---------------------------------------------------------------------------------------

MIN_CHARS_PER_PAGE = 20  # a page with less text than this is treated as an image page


def extract_pdf(data: bytes) -> Extracted:
    check_size(data)
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise Unscannable(UnscannableReason.ENCRYPTED)
        out = Extracted()
        empty_pages = 0
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if len(text.strip()) < MIN_CHARS_PER_PAGE:
                empty_pages += 1
            for line in text.splitlines():
                out.add_text(line, f"{i}쪽")
    except Unscannable:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError):
        raise Unscannable(UnscannableReason.CORRUPT) from None
    if empty_pages >= len(reader.pages):
        raise Unscannable(UnscannableReason.IMAGE_ONLY)  # no page has text: scanned PDF
    out.partial = empty_pages > 0  # some pages are images: scanned only in part
    return out


# --- HWPX (OWPML) ------------------------------------------------------------------------------

_SECTION_RE = re.compile(r"^Contents/section(\d+)\.xml$", re.IGNORECASE)


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def extract_hwpx(data: bytes) -> Extracted:
    check_size(data)
    zf = open_zip_safely(data)
    from defusedxml import ElementTree
    from defusedxml.common import DefusedXmlException

    try:
        names = zf.namelist()
        manifest = next((n for n in names if n.lower() == "meta-inf/manifest.xml"), None)
        if manifest is not None and b"encryption-data" in zf.read(manifest):
            raise Unscannable(UnscannableReason.ENCRYPTED)
        sections = sorted((int(m.group(1)), n) for n in names if (m := _SECTION_RE.match(n)))
        if not sections:
            raise Unscannable(UnscannableReason.CORRUPT)
        out = Extracted()
        counters = {"p": 0, "tbl": 0}
        for _idx, name in sections:
            _walk_hwpx(ElementTree.fromstring(zf.read(name)), out, counters)
        return out
    except Unscannable:
        raise
    except (DefusedXmlException, ElementTree.ParseError, KeyError, ValueError):
        raise Unscannable(UnscannableReason.CORRUPT) from None
    finally:
        zf.close()


def _walk_hwpx(elem: Any, out: Extracted, counters: dict[str, int]) -> None:
    """Paragraphs outside tables become segments; tables become Table rows.

    In OWPML a table sits inside a paragraph run (<p><run><tbl>…), so paragraph text excludes
    table subtrees and the tables are collected separately.
    """
    tag = _local(elem.tag)
    if tag == "tbl":
        counters["tbl"] += 1
        rows = [
            [_cell_text(tc) for tc in tr if _local(tc.tag) == "tc"]
            for tr in elem
            if _local(tr.tag) == "tr"
        ]
        out.tables.append(Table(f"표 {counters['tbl']}", rows[:MAX_TABLE_ROWS]))
        return
    if tag == "p":
        counters["p"] += 1
        out.add_text(_text_excluding_tables(elem), f"문단 {counters['p']}")
    for child in elem:
        _walk_hwpx(child, out, counters)


def _text_excluding_tables(elem: Any) -> str:
    parts: list[str] = []

    def rec(e: Any) -> None:
        tag = _local(e.tag)
        if tag == "tbl":
            return
        if tag == "t" and e.text:
            parts.append(e.text)
        for c in e:
            rec(c)

    rec(elem)
    return "".join(parts)


def _cell_text(tc: Any) -> str:
    texts = [_text_excluding_tables(p) for p in tc.iter() if _local(p.tag) == "p"]
    return " ".join(t for t in texts if t).strip()
